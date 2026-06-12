#!/usr/bin/env python3
"""把旧的单表 dish.xlsx 拆成 dish_base.xlsx（本体）+ dish_variant.xlsx（变体/库条目）。

拆分规则：
- 本体表保留物理属性：id/name/deliciousness/icon/allowRotate + 多行 *shapeRows。
- 变体表是真正的随机库条目：id(沿用旧 dishId，保证 recipe 引用不变) / baseId(=旧 id) /
  aTagId / bTagId（唯一标签槽位，初始为空） / hiddenMin/Max / baseWeight / price / maxRollCount
  + 多行 *inherentTags（固有标签）。
- price 为新增字段，旧数据没有，这里按 id 给出占位价格（见 PRICE）。

用法：
    python3 GameConfig/Tools/split_dish_tables.py
依赖：openpyxl。
"""
import os

from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
DATAS = os.path.normpath(os.path.join(HERE, "..", "Datas"))

# 占位价格（设计文档要求 price 字段；旧数据无，按菜品手动给一档）。
PRICE = {
    "rice": 2, "egg": 3, "soup": 3, "fish": 5, "noodle": 4,
    "dumpling": 4, "chili": 2, "orange": 3, "cake": 6, "feast": 8,
}
DEFAULT_PRICE = 3


def read_dishes():
    """从旧 dish.xlsx 读出每条记录，处理 *shapeRows 多行续行。"""
    wb = load_workbook(os.path.join(DATAS, "dish.xlsx"), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [c for c in rows[0]]  # ["##var", "id", ... , "*shapeRows"]
    fields = [(c[1:] if isinstance(c, str) and c.startswith("*") else c) for c in header]
    # 字段名 -> 列下标（跳过列 0 的 ##var 标记）。
    idx = {name: i for i, name in enumerate(fields) if name not in (None, "##var")}

    dishes = []
    for row in rows[2:]:  # 跳过 ##var / ##type 两行
        id_cell = row[idx["id"]]
        shape_cell = row[idx["shapeRows"]]
        if id_cell:  # 新记录起始行
            rec = {name: row[i] for name, i in idx.items()}
            rec["shapeRows"] = [shape_cell] if shape_cell else []
            dishes.append(rec)
        elif dishes and shape_cell:  # 续行：只补 shapeRows
            dishes[-1]["shapeRows"].append(shape_cell)
    return dishes


def b(v):
    return "true" if v in (True, "true", "True", 1) else "false"


def write_base(dishes):
    wb = Workbook()
    ws = wb.active
    ws.title = "dish_base"
    ws.append(["##var", "id", "name", "deliciousness", "icon", "allowRotate", "*shapeRows"])
    ws.append(["##type", "string", "string", "int", "string", "bool", "list,string"])
    for d in dishes:
        shapes = d["shapeRows"] or [""]
        base = ["", d["id"], d["name"], d["deliciousness"], d["icon"], b(d["allowRotate"])]
        ws.append(base + [shapes[0]])
        for s in shapes[1:]:
            ws.append([""] * len(base) + [s])
    wb.save(os.path.join(DATAS, "dish_base.xlsx"))
    print(f"  dish_base.xlsx ({len(dishes)} 本体)")


def write_variant(dishes):
    wb = Workbook()
    ws = wb.active
    ws.title = "dish_variant"
    ws.append(["##var", "id", "baseId", "aTagId", "bTagId", "hiddenMin", "hiddenMax",
               "baseWeight", "price", "maxRollCount", "*inherentTags"])
    ws.append(["##type", "string", "string", "string", "string", "int", "int",
               "float", "int", "int", "list,string"])
    for d in dishes:
        tags = [t for t in (d.get("inherentTags"),) if t]  # 旧表 inherentTags 为单格
        if isinstance(d.get("inherentTags"), str) and d["inherentTags"]:
            tags = [d["inherentTags"]]
        else:
            tags = []
        price = PRICE.get(d["id"], DEFAULT_PRICE)
        base = ["", d["id"], d["id"], "", "", d["hiddenMin"], d["hiddenMax"],
                d["baseWeight"], price, d["maxRollCount"]]
        ws.append(base + [tags[0] if tags else ""])
        for t in tags[1:]:
            ws.append([""] * len(base) + [t])
    wb.save(os.path.join(DATAS, "dish_variant.xlsx"))
    print(f"  dish_variant.xlsx ({len(dishes)} 基础变体)")


def main():
    dishes = read_dishes()
    print(f"读到 {len(dishes)} 条菜品，开始拆表 -> {DATAS}")
    write_base(dishes)
    write_variant(dishes)
    print("拆表完成。请运行 bash GameConfig/gen.sh，并删除旧 dish.xlsx。")


if __name__ == "__main__":
    main()
