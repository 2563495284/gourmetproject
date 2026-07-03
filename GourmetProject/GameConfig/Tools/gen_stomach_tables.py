#!/usr/bin/env python3
"""生成「胃部棋盘 + 格子标签」相关的 xlsx 表：
- stomach_fragment.xlsx：胃部碎片库（含初始胃 gut_4x4）。
- fragment_cell_tag.xlsx：碎片格强化标签（分开配置）。
- 重写 character.xlsx：补齐 initialFragmentId / maxStomachWidth / maxStomachHeight / timelinePool / bossPool。
- 重写 tag.xlsx：追加 2 个「格子强化」标签 t_cell_gold / t_cell_warm。

用法：python3 GameConfig/Tools/gen_stomach_tables.py
依赖：openpyxl。
"""
import os

from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
DATAS = os.path.normpath(os.path.join(HERE, "..", "Datas"))

# —— 碎片库 —— (id, hiddenMin, hiddenMax, baseWeight, price, shapeRows)
FRAGMENTS = [
    ("gut_4x4", 0, 0, 0, 0, ["XXXX", "XXXX", "XXXX", "XXXX"]),       # 初始胃，不入随机池
    ("frag_1x2", 1, 10, 100, 2, ["XX"]),
    ("frag_2x1", 1, 10, 100, 2, ["X", "X"]),                          # 与 1x2 不同碎片(不旋转)
    ("frag_L", 10, 25, 70, 4, ["XX", ".X"]),
    ("frag_gold2x2", 25, 45, 40, 6, ["XX", "XX"]),                    # (0,0) 带强化格
]

# —— 碎片格强化标签 —— (id, fragmentId, x, y, tagId)
CELL_TAGS = [
    ("fct_gold2x2_0_0", "frag_gold2x2", 0, 0, "t_cell_gold"),
    ("fct_L_1_1", "frag_L", 1, 1, "t_cell_warm"),
]

# —— 追加到 tag.xlsx 的格子强化标签 —— (与现有列顺序一致)
NEW_TAGS = [
    ("t_cell_gold", "强化·黄金格", "占据该格的菜品贡献分数 ×1.5。", "Inherent", "AddMult", "1.5", "", ""),
    ("t_cell_warm", "强化·温补格", "占据该格的菜品美味度 +3。", "Inherent", "AddFlat", "3", "", ""),
]


def write_fragments():
    # hiddenRange 用 HiddenRange bean(单元格写 min,max)，放在多行列 *shapeRows 之前。
    wb = Workbook()
    ws = wb.active
    ws.title = "stomach_fragment"
    ws.append(["##var", "id", "baseWeight", "price", "hiddenRange", "*shapeRows"])
    ws.append(["##type", "string", "float", "int", "HiddenRange", "list,string"])
    for fid, hmin, hmax, w, price, rows in FRAGMENTS:
        ws.append(["", fid, w, price, f"{hmin},{hmax}", rows[0]])
        for r in rows[1:]:
            ws.append(["", "", "", "", "", r])
    wb.save(os.path.join(DATAS, "stomach_fragment.xlsx"))
    print(f"  stomach_fragment.xlsx ({len(FRAGMENTS)} 碎片)")


def write_cell_tags():
    wb = Workbook()
    ws = wb.active
    ws.title = "fragment_cell_tag"
    ws.append(["##var", "id", "fragmentId", "x", "y", "tagId"])
    ws.append(["##type", "string", "string", "int", "int", "string"])
    for row in CELL_TAGS:
        ws.append([""] + list(row))
    wb.save(os.path.join(DATAS, "fragment_cell_tag.xlsx"))
    print(f"  fragment_cell_tag.xlsx ({len(CELL_TAGS)} 格标签)")


def rewrite_characters():
    src = os.path.join(DATAS, "character.xlsx")
    rows = list(load_workbook(src, data_only=True).active.iter_rows(values_only=True))
    header = next(row for row in rows if row and row[0] == "##var")
    index = {name: i for i, name in enumerate(header) if name}
    data = [row for row in rows if row and row[0] not in ("##var", "##comment", "##type")]
    wb = Workbook()
    ws = wb.active
    ws.title = "character"
    ws.append(["##var", "id", "name", "desc", "portrait", "initialRecipeId",
               "initialFragmentId", "maxStomachWidth", "maxStomachHeight",
               "timelinePool", "bossPool", "initialGold", "startItems"])
    ws.append(["##comment", "配置ID", "角色名称", "角色描述", "角色立绘资源路径", "初始菜谱ID",
               "初始胃碎片ID", "胃最大宽度", "胃最大高度",
               "可用行动轴池，空=全部，逗号分隔 timeline.id",
               "可用Boss池，空=全部，逗号分隔 boss.id", "初始金币", "初始携带道具ID列表"])
    ws.append(["##type", "string", "string", "string", "string", "string",
               "string", "int", "int", "string", "string", "int", "list,string"])
    for r in data:
        def value(name, default=""):
            i = index.get(name)
            if i is None or i >= len(r) or r[i] is None:
                return default
            return r[i]

        ws.append([
            "",
            value("id"),
            value("name"),
            value("desc"),
            value("portrait"),
            value("initialRecipeId"),
            value("initialFragmentId", "gut_4x4"),
            value("maxStomachWidth", 4),
            value("maxStomachHeight", 4),
            value("timelinePool"),
            value("bossPool"),
            value("initialGold", 9999),
            value("startItems"),
        ])
    wb.save(src)
    print(f"  character.xlsx 重写/补齐字段 ({len(data)} 角色)")


def append_tags():
    src = os.path.join(DATAS, "tag.xlsx")
    wb = load_workbook(src)
    ws = wb.active
    existing = {row[1] for row in ws.iter_rows(min_row=3, values_only=True) if row[1]}
    added = 0
    for t in NEW_TAGS:
        if t[0] in existing:
            continue
        ws.append([""] + list(t))
        added += 1
    wb.save(src)
    print(f"  tag.xlsx 追加 {added} 个格子强化标签")


def main():
    print(f"Datas 目录: {DATAS}")
    write_fragments()
    write_cell_tags()
    rewrite_characters()
    append_tags()
    print("生成完成。请运行 bash GameConfig/gen.sh。")


if __name__ == "__main__":
    main()
