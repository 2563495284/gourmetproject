#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把「道具设计」sheet 的设计条目落地为 item.xlsx `item` 表的正式配置行。

- 保留已有 5 条道具（item_big_plate/chef_knife/pepper_jar/reroll/extra_serve）不动。
- 幂等：重复运行会先删除本脚本生成的 item_* 行（除已有 5 条外）再重写。
- 用 ws.cell 显式写入，避免 openpyxl append 破坏空单元格。
效果类型字符串须与 Assets/.../Game/Meta/Items/ItemEffectTypes.cs 常量一致。
"""
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "..", "Datas", "item.xlsx")

KEEP_IDS = {
    "item_big_plate", "item_chef_knife", "item_pepper_jar",
    "item_reroll", "item_extra_serve",
}

# (id, name, desc, kind, quality, specialTags, triggerTiming, holdLimit, effectType, effectValue, effectParam)
P, A = "Passive", "Active"
ROWS = [
    # 通用正面 - 随机获得（一次性主动，进入选择/发放流程）
    ("item_grant_two_passive", "随机被动包", "获得两个随机被动道具。", A, "Uncommon", "", "RewardScreen", 9, "GrantRandomPassive", 2, ""),
    ("item_choose_one_passive", "被动三选一", "从多个被动道具中选择一个获得。", A, "Uncommon", "", "RewardScreen", 9, "ChooseOnePassive", 3, ""),
    ("item_grant_two_active", "随机主动包", "获得两个随机主动道具。", A, "Uncommon", "", "RewardScreen", 9, "GrantRandomActive", 2, ""),
    ("item_choose_one_active", "主动三选一", "从多个主动道具中选择一个获得。", A, "Uncommon", "", "RewardScreen", 9, "ChooseOneActive", 3, ""),
    ("item_choose_one_food", "食物三选一", "从多个食物中选择一个获得。", A, "Common", "", "RewardScreen", 9, "ChooseOneFood", 3, ""),
    ("item_choose_one_fragment", "碎片三选一", "从多个胃部碎片中选择一个获得。", A, "Uncommon", "", "RewardScreen", 9, "ChooseOneFragment", 3, ""),
    ("item_grant_recipe", "菜谱券", "获得一个菜谱。", A, "Uncommon", "", "RewardScreen", 9, "GrantRecipe", 1, ""),
    ("item_family_pack", "全家福", "获得一个随机被动道具、随机食物、50金币。", A, "Rare", "", "RewardScreen", 9, "FamilyPack", 50, ""),
    ("item_randomize_items", "混沌重置", "所有道具替换为随机道具。", A, "Rare", "", "RewardScreen", 9, "RandomizeItems", 0, ""),
    # 给钱
    ("item_gold_random", "随机金币", "获得 1~100 个金币。", A, "Common", "", "RewardScreen", 9, "GoldNow", 0, "range:1,100"),
    ("item_gold_meal_bonus", "美食分红", "接下来10局美食，额外获得15金币。", P, "Uncommon", "", "None", 1, "GoldMealBonus", 15, "meals:10"),
    ("item_gold_boss", "Boss赏金", "完成本次 Boss 后，额外获得100金币。", P, "Uncommon", "", "None", 1, "GoldOnBossComplete", 100, ""),
    ("item_gold_percent", "利润提成", "美食结束获得的金币，额外获得20%。", P, "Rare", "", "None", 1, "GoldMealPercent", 0.2, ""),
    # 打折
    ("item_discount_food", "食材优惠券", "购买食物打折20%。", P, "Common", "", "None", 1, "ShopDiscountFood", 0.2, ""),
    ("item_discount_fragment", "碎片优惠券", "购买碎片打折20%。", P, "Common", "", "None", 1, "ShopDiscountFragment", 0.2, ""),
    ("item_discount_active", "主动道具券", "购买主动道具打折20%。", P, "Common", "", "None", 1, "ShopDiscountActive", 0.2, ""),
    ("item_discount_passive", "被动道具券", "购买被动道具打折20%。", P, "Common", "", "None", 1, "ShopDiscountPassive", 0.2, ""),
    ("item_discount_remove", "删牌优惠", "删牌价格打折20%。", P, "Common", "", "None", 1, "ShopDiscountRemove", 0.2, ""),
    ("item_remove_fixed_30", "固定删牌价", "删牌价格固定为30金币。", P, "Uncommon", "", "None", 1, "RemovePriceFixed", 30, ""),
    ("item_discount_recipe", "菜谱优惠", "购买菜谱打折20%。", P, "Common", "", "None", 1, "ShopDiscountRecipe", 0.2, ""),
    # 不死
    ("item_famous_knife", "名刀·加护", "分数未达标时不失败，但会失去此道具。", P, "Rare", "", "None", 1, "Undying", 0, ""),
    # 分数降低
    ("item_score_to_one", "免检通行证", "接下来3局非盛宴，分数要求为1。", A, "Rare", "", "RewardScreen", 9, "RequiredScoreToOne", 0, "meals:3"),
    ("item_req_normal_down", "家常减负", "普通美食要求分数-10%。", P, "Uncommon", "", "None", 1, "RequiredScoreNormalPct", -0.1, ""),
    ("item_req_super_down", "超级减负", "超级美食要求分数-10%。", P, "Uncommon", "", "None", 1, "RequiredScoreSuperPct", -0.1, ""),
    ("item_req_feast_down", "盛宴减负", "盛宴美食要求分数-10%。", P, "Uncommon", "", "None", 1, "RequiredScoreFeastPct", -0.1, ""),
    # 数量检测
    ("item_count_le_mult", "精致料理", "结算时食物数量<=10，所有美食倍率×1.1。", P, "Uncommon", "", "None", 1, "CountThresholdFinalMult", 1.1, "lte:10"),
    ("item_count_ge_mult", "丰盛大餐", "结算时食物数量>=10，所有美食倍率×1.1。", P, "Uncommon", "", "None", 1, "CountThresholdFinalMult", 1.1, "gte:10"),
    # 结算次数
    ("item_per_dish_mult", "流水席", "结算时每结算1个食物，所有食物倍率+0.1。", P, "Rare", "", "None", 1, "PerDishSettledMultFlat", 0.1, ""),
    # 分数/倍率增加
    ("item_perma_flat_all", "增鲜剂", "所有食物基础分数永久+10。", P, "Rare", "", "None", 1, "PermanentAddFlatAll", 10, ""),
    ("item_perma_mult_all", "浓缩精华", "所有食物基础倍率永久+0.2。", P, "Rare", "", "None", 1, "PermanentAddMultAll", 0.2, ""),
    # 上菜
    ("item_stargaze_every5", "观星仪", "每5次上菜后，下一次上菜可预见食物。", P, "Uncommon", "", "None", 1, "StarGazeEveryN", 0, "every:5"),
    ("item_stargaze_first3", "开局观星", "每局前三次上菜可预见食物。", P, "Uncommon", "", "None", 1, "StarGazeFirstN", 0, "first:3"),
    ("item_free_move_first", "免费调整", "每局第一次上菜后，可不消耗修正次数移动。", P, "Uncommon", "", "None", 1, "FreeMoveFirst", 0, ""),
    # 调整
    ("item_adjust_plus1", "熟练的手", "调整次数+1。", P, "Common", "", "None", 1, "AdjustCountBonus", 1, ""),
    ("item_adjust_to_mult", "节俭加成", "每1个未使用的调整次数，使倍率+0.2。", P, "Rare", "", "None", 1, "AdjustToMult", 0.2, ""),
    # 食物转换
    ("item_food_convert", "点石成金", "随机将5个食物变为至少同品质的食物。", A, "Rare", "", "BeforeEat", 9, "FoodConvert", 5, ""),
    # 运营
    ("item_adjust_to_gold", "变卖技巧", "每个未使用的调整次数，获得1金币。", P, "Uncommon", "", "None", 1, "GoldPerUnusedAdjust", 1, ""),
    ("item_gold_on_shop", "商会返利", "进入商店时，获得10金币。", P, "Uncommon", "", "None", 1, "GoldOnShopEnter", 10, ""),
    ("item_extra_interest", "复利账户", "每周结束额外执行一次利息。", P, "Rare", "", "None", 1, "ExtraInterest", 1, ""),
    ("item_min_gold", "保底基金", "回合结束金币<10时，补足到10金币。", P, "Uncommon", "", "None", 1, "MinGoldGuarantee", 10, ""),
    ("item_interest_cap", "高息账户", "利息上限扩展到10金币。", P, "Rare", "", "None", 1, "InterestCapBonus", 10, ""),
    ("item_loan", "高利贷", "临时获得300金币，下一周失去600金币。", A, "Rare", "", "RewardScreen", 9, "Loan", 300, "repay:600"),
    # 时间轴
    ("item_timeline_random", "命运骰子", "随机化事件轴的节点时间。", A, "Uncommon", "", "RewardScreen", 9, "TimelineRandomize", 0, ""),
    ("item_extra_day", "加班一天", "延后所有 Boss 一天。", A, "Rare", "", "RewardScreen", 9, "TimelineExtraDay", 1, ""),
    ("item_week_minus", "时光倒流", "周数-1。", A, "Epic", "", "RewardScreen", 9, "TimelineWeekMinus", 1, ""),
    ("item_add_reward_node", "额外奖励", "增加一个奖励节点到时间轴上。", A, "Uncommon", "", "RewardScreen", 9, "TimelineAddRewardNode", 1, ""),
    ("item_add_interest_node", "额外利息节点", "增加一个利息节点到时间轴上。", A, "Uncommon", "", "RewardScreen", 9, "TimelineAddInterestNode", 1, ""),
    # 分数加倍
    ("item_every3_next_mult", "节奏大师", "每上3个食物后，下一个食物倍率+2。", P, "Rare", "", "None", 1, "EveryNthServeMult", 2, "every:3"),
    ("item_first_x2", "开门红", "第一个上的食物，倍率×2。", P, "Uncommon", "", "None", 1, "NthServeMult", 2, "index:1"),
    ("item_last_x2", "压轴好戏", "最后一个上的食物，倍率×2。", P, "Uncommon", "", "None", 1, "NthServeMult", 2, "index:-1"),
    # 主动道具
    ("item_extra_active_slots", "工具箱", "增加2个额外主动道具栏。", P, "Rare", "", "None", 1, "ExtraActiveSlot", 2, ""),
    ("item_gold_on_active", "以用换钱", "使用主动道具时，获得20金币。", P, "Uncommon", "", "None", 1, "GoldOnActiveUse", 20, ""),
    ("item_block_active", "极简主义", "无法获得主动道具，立即获得1000金币。", P, "Epic", "", "None", 1, "BlockActive", 1000, ""),
    # 风味
    ("item_flavor_enhance", "风味喷雾", "强化两个随机风味到无风味食物上。", A, "Uncommon", "", "BeforeEat", 9, "FlavorEnhance", 2, ""),
    ("item_flavor_remove_gold", "风味回收", "移除一个食物的风味，获得200金币。", A, "Uncommon", "", "BeforeEat", 9, "FlavorRemoveForGold", 200, ""),
    ("item_flavor_remove_copyskill", "风味转化", "移除一个食物的风味，其技能复制。", A, "Rare", "", "BeforeEat", 9, "FlavorRemoveCopySkill", 0, ""),
    ("item_flavor_remove_double", "风味浓缩", "移除一个食物的风味，其分数翻倍。", A, "Rare", "", "BeforeEat", 9, "FlavorRemoveDoubleScore", 2, ""),
    ("item_flavor_double_slot", "双重风味", "允许一个食物拥有两个风味。", P, "Epic", "", "None", 1, "FlavorDoubleSlot", 1, ""),
    ("item_flavor_contagion", "风味传染", "将一个食物的风味强化到无风味食物上。", A, "Rare", "", "BeforeEat", 9, "FlavorContagion", 1, ""),
    # 标签
    ("item_celltag_enhance", "标签强化", "强化两个随机标签到胃部格子上。", A, "Uncommon", "", "BeforeEat", 9, "CellTagEnhance", 2, ""),
    ("item_celltag_contagion", "标签传染", "将一个格子的标签强化到其他格子上。", A, "Rare", "", "BeforeEat", 9, "CellTagContagion", 1, ""),
    # 行动
    ("item_lucky_guarantee", "好运连连", "每4个事件的最后一个是奖励事件。", P, "Rare", "", "None", 1, "LuckyEventGuarantee", 4, ""),
    ("item_lucky_chance", "幸运符", "事件中遇到奖励的概率提高。", P, "Uncommon", "", "None", 1, "LuckyEventChance", 0.2, ""),
    ("item_more_events", "热闹街区", "遇到事件的概率提高。", P, "Uncommon", "", "None", 1, "MoreEvents", 0.2, ""),
    ("item_gold_on_event", "事件红包", "完成一个事件，获得10金币。", P, "Common", "", "None", 1, "GoldOnEventComplete", 10, ""),
    ("item_reroll_action", "重掷行动", "获得2次重新随机行动的机会。", A, "Common", "", "RewardScreen", 9, "RerollAction", 2, ""),
    # 商店
    ("item_shop_restock", "自动补货", "商店会自动补货。", P, "Rare", "", "None", 1, "ShopRestock", 0, ""),
    # 美食
    ("item_choice_count_plus1", "琳琅满目", "多选一食物时，可选数量+1。", P, "Uncommon", "", "None", 1, "ChoiceCountBonus", 1, ""),
    ("item_choice_times_plus1", "细嚼慢咽", "多选一食物时，可选次数+1，可选数量-1。", P, "Rare", "", "None", 1, "ChoiceTimesBonus", 1, ""),
    ("item_extra_food_choice", "加餐", "每3次普通美食，额外获得一个多选一美食。", P, "Rare", "", "None", 1, "ExtraFoodChoice", 0, "every:3"),
    ("item_extra_item_choice", "加料", "每3次超级美食，额外获得一个多选一物品。", P, "Rare", "", "None", 1, "ExtraItemChoice", 0, "every:3"),
    ("item_copy_food", "复制美食", "随机复制一个美食。", A, "Rare", "", "RewardScreen", 9, "CopyFood", 1, ""),
    # 负面处理
    ("item_discard_negative", "断舍离", "丢弃两个负面道具。", A, "Common", "", "RewardScreen", 9, "DiscardNegative", 2, ""),
    ("item_discard_negative_gold", "变废为宝", "丢弃所有负面道具，每丢弃一个获得100金币。", A, "Uncommon", "", "RewardScreen", 9, "DiscardNegativeForGold", 100, ""),

    # 通用负面
    ("item_gold_meal_penalty", "克扣工钱", "美食获得的金币减少25%。", P, "Common", "Negative", "None", 1, "GoldMealPercent", -0.25, ""),
    ("item_shop_price_up", "涨价", "商店的价格变高25%。", P, "Common", "Negative", "None", 1, "ShopPriceUp", 0.25, ""),
    ("item_no_remove", "囤积癖", "无法再删除食物。", P, "Common", "Negative", "None", 1, "NoRemoveDish", 0, ""),
    ("item_skip_node", "停业整顿", "跳过下一个商店或利息节点。", P, "Common", "Negative", "None", 1, "TimelineSkipNode", 1, ""),
    ("item_choice_minus1", "选择困难", "多选一食物时，可选数量-1。", P, "Common", "Negative", "None", 1, "ChoiceCountPenalty", -1, ""),
    ("item_req_normal_up", "家常加码", "普通美食要求分数+10%。", P, "Common", "Negative", "None", 1, "RequiredScoreNormalPct", 0.1, ""),
    ("item_req_super_up", "超级加码", "超级美食要求分数+10%。", P, "Common", "Negative", "None", 1, "RequiredScoreSuperPct", 0.1, ""),
    ("item_req_feast_up", "盛宴加码", "盛宴美食要求分数+10%。", P, "Common", "Negative", "None", 1, "RequiredScoreFeastPct", 0.1, ""),
    ("item_gold_week_clear", "月光族", "本周结束时，金币清空。", P, "Common", "Negative", "None", 1, "GoldWeekClear", 0, ""),

    # 专有
    ("item_count_as_all", "小份主义", "所有食物额外被视为1个食物。", P, "Rare", "", "None", 1, "CountAsBonusAll", 1, ""),
    ("item_transfer_target_mult", "传递受益", "传递时，被传递方倍率×1.1。", P, "Rare", "", "None", 1, "TransferTargetMult", 1.1, ""),
    ("item_transfer_source_mult", "传递奉献", "传递时，传递方倍率×1.1。", P, "Rare", "", "None", 1, "TransferSourceMult", 1.1, ""),
    ("item_skill_count_mult", "技法大师", "结算时每一个技能，所有食物倍率+0.1。", P, "Rare", "", "None", 1, "PerSkillMultFlat", 0.1, ""),
    ("item_gold_on_transfer", "传递分红", "传递10次后，获得10金币。", P, "Uncommon", "", "None", 1, "GoldOnTransferCount", 10, "count:10"),
    ("item_cake_retain", "蛋糕保鲜", "每次保留10%的蛋糕层数至下一次美食。", P, "Rare", "", "None", 1, "CakeLayerRetain", 0.1, ""),
    ("item_cake_init_bonus", "蛋糕打底", "蛋糕层数初始额外获得10层。", P, "Uncommon", "", "None", 1, "CakeLayerInitBonus", 10, ""),
    ("item_cake_accel", "蛋糕膨胀", "蛋糕层数增加时，额外加1层。", P, "Rare", "", "None", 1, "CakeLayerAccel", 1, ""),
    ("item_cake_to_gold", "蛋糕变现", "蛋糕层数大于100层时，获得10金币。", P, "Uncommon", "", "None", 1, "GoldOnCakeLayers", 10, "threshold:100"),
    ("item_cake_req_minus", "蛋糕捷径", "蛋糕效果需求的层数-10。", P, "Uncommon", "", "None", 1, "CakeLayerReqMinus", 10, ""),
]

HIDDEN_BY_QUALITY = {
    "Common": "1,30",
    "Uncommon": "5,45",
    "Rare": "20,65",
    "Epic": "40,85",
    "Legendary": "60,100",
}


def main():
    path = os.path.abspath(XLSX)
    wb = openpyxl.load_workbook(path)
    ws = wb["item"]

    # 删除本脚本产出的旧行（保留表头与已有 5 条），实现幂等重写。
    keep_rows = []
    for r in range(4, ws.max_row + 1):
        rid = ws.cell(row=r, column=2).value
        if rid in KEEP_IDS:
            keep_rows.append(r)

    # 找到已有 5 条之后的第一空行；从那里开始写。
    start = max(keep_rows) + 1 if keep_rows else 4
    # 清空 start 以后的所有单元格（O 列 = 15）。
    for r in range(start, ws.max_row + 1):
        for c in range(1, 16):
            ws.cell(row=r, column=c, value=None)

    seen = set()
    row = start
    for (rid, name, desc, kind, quality, tags, timing, hold, etype, eval_, eparam) in ROWS:
        assert rid not in seen and rid not in KEEP_IDS, f"duplicate id {rid}"
        seen.add(rid)
        lwp = "3,0.65,1" if kind == "Passive" else "1,1,1"
        hidden = HIDDEN_BY_QUALITY.get(quality, "1,30") if kind == "Passive" else "0,0"
        vals = [None, rid, name, desc, kind, quality, tags, timing, hold, etype, eval_, eparam, "", lwp, hidden]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=row, column=c, value=v)
        row += 1

    wb.save(path)
    print(f"wrote {len(ROWS)} item rows into {path} (rows {start}..{row - 1})")


if __name__ == "__main__":
    main()
