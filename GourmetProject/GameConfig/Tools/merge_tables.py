#!/usr/bin/env python3
"""把同主题的多张单表 xlsx 合并成一个多 sheet 文件，并同步 __tables__.xlsx 的 input。

Luban 的 sheet 指定语法是 `sheet名@文件.xlsx`（sheet 名在前），且 sheet 的 A1
单元格必须以 `##` 开头才会被当数据表读取。本项目各数据 xlsx 的 A1 都是 `##var`，天然满足。

幂等：源既可能是「旧的单表文件」，也可能是「已经合并好的多 sheet 文件里的某个 sheet」，
两种状态都能正确读出并重写，所以脚本可反复运行。

用法：
    python3 GameConfig/Tools/merge_tables.py
依赖：openpyxl。
"""
import os

from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
DATAS = os.path.normpath(os.path.join(HERE, "..", "Datas"))

# 目标文件 -> [(sheet 名, 旧单表文件名)]。sheet 名同时用作 __tables__ 的 input 前缀。
GROUPS = {
    "dish.xlsx": [
        ("dish_base", "dish_base.xlsx"),
        ("dish_variant", "dish_variant.xlsx"),
    ],
    "event.xlsx": [
        ("event", "event.xlsx"),
        ("event_option", "event_option.xlsx"),
    ],
    "timeline.xlsx": [
        ("timeline", "timeline.xlsx"),
        ("timeline_node", "timeline_node.xlsx"),
    ],
    "reward.xlsx": [
        ("reward_package", "reward_package.xlsx"),
        ("reward_slot", "reward_slot.xlsx"),
        ("reward_pool", "reward_pool.xlsx"),
    ],
}

# full_name -> input。sheet@文件 形式精确指定数据 sheet。
INPUT_MAP = {
    "TbDishBase": "dish_base@dish.xlsx",
    "TbDishVariant": "dish_variant@dish.xlsx",
    "TbEvent": "event@event.xlsx",
    "TbEventOption": "event_option@event.xlsx",
    "TbTimeline": "timeline@timeline.xlsx",
    "TbTimelineNode": "timeline_node@timeline.xlsx",
    "TbRewardPackage": "reward_package@reward.xlsx",
    "TbRewardSlot": "reward_slot@reward.xlsx",
    "TbRewardPool": "reward_pool@reward.xlsx",
}


def read_sheet_rows(target, sheet, src):
    """读出某张表的所有行。优先读旧单表，其次读已合并多 sheet 文件里的对应 sheet。"""
    src_path = os.path.join(DATAS, src)
    if os.path.exists(src_path):
        wb = load_workbook(src_path, data_only=True)
        ws = wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    tgt_path = os.path.join(DATAS, target)
    if os.path.exists(tgt_path):
        wb = load_workbook(tgt_path, data_only=True)
        if sheet in wb.sheetnames:
            return [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    raise FileNotFoundError(f"找不到 {sheet} 的数据源（{src} 或 {target}@{sheet}）")


def main():
    # 1) 先把全部数据读进内存，避免重写 event.xlsx/timeline.xlsx 时覆盖掉自身源。
    data = {}
    srcfiles = set()
    for target, members in GROUPS.items():
        for sheet, src in members:
            data[(target, sheet)] = read_sheet_rows(target, sheet, src)
            srcfiles.add(src)

    # 2) 写多 sheet 目标文件。
    for target, members in GROUPS.items():
        wb = Workbook()
        first = True
        for sheet, _src in members:
            ws = wb.active if first else wb.create_sheet(title=sheet)
            if first:
                ws.title = sheet
                first = False
            for row in data[(target, sheet)]:
                ws.append(row)
        wb.save(os.path.join(DATAS, target))
        print(f"  写出 {target} {wb.sheetnames}")

    # 3) 删除已被并入多 sheet、且不再作为目标名存在的旧单表。
    keep = set(GROUPS.keys())
    for src in sorted(srcfiles):
        if src not in keep:
            p = os.path.join(DATAS, src)
            if os.path.exists(p):
                os.remove(p)
                print(f"  删除旧表 {src}")

    # 4) 同步 __tables__.xlsx 的 input 列（B=full_name 第2列，E=input 第5列）。
    tables_path = os.path.join(DATAS, "__tables__.xlsx")
    wb = load_workbook(tables_path)
    ws = wb.active
    for row in ws.iter_rows():
        fn = row[1].value
        if fn in INPUT_MAP:
            row[4].value = INPUT_MAP[fn]
            print(f"  登记 {fn} -> {INPUT_MAP[fn]}")
    wb.save(tables_path)

    print("合并完成。请运行 bash GameConfig/gen.sh 重新生成。")


if __name__ == "__main__":
    main()
